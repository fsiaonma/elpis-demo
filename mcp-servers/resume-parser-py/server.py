#!/usr/bin/env python3
"""Resume parser MCP server — stdin/stdout JSON-RPC, parse_resume tool only."""

from __future__ import annotations

import json
import re
import sys
import time
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document

LINE_TOP_TOLERANCE = 4.0
SERVER_DIR = Path(__file__).resolve().parent
RESUMES_ROOT = (SERVER_DIR / "../../data/resumes").resolve()

SECTION_KEYWORDS: dict[str, list[str]] = {
    "education": ["教育背景", "教育经历", "学历"],
    "experience": ["工作经历", "工作经验", "职业经历"],
    "projects": ["项目经验", "项目经历"],
    "skills": ["专业技能", "技术专长", "技能特长", "技术栈", "技能"],
}

SHORT_SECTION_KEYWORDS: dict[str, str] = {
    "工作": "experience",
    "教育": "education",
    "项目": "projects",
    "技能": "skills",
}

DATE_RANGE_RE = re.compile(
    r"(?P<start>\d{4}\.\d{2})\s*[-–—]\s*(?P<end>至今|\d{4}\.\d{2})"
)
PHONE_RE = re.compile(r"1\d{10}")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.\w+")
NAME_RE = re.compile(
    r"(?:姓名|Name)\s*[：:]\s*"
    r"(?P<name>[^\s|，,;；|]+(?:\s+[^\s|，,;；|]+)?)"
    r"(?=\s*(?:学历|电话|邮箱|Email|Phone|求职|性别|年龄|地址|Location|$))",
    re.IGNORECASE,
)
DEGREE_INLINE_RE = re.compile(
    r"学历\s*[：:]\s*(?P<degree>本科|硕士|博士|大专)"
    r"(?:\s*[,，;；|]\s*专业\s*[：:]\s*(?P<major>[^\s|，,;；]+))?"
)
MAJOR_INLINE_RE = re.compile(r"专业\s*[：:]\s*(?P<major>[^\s|，,;；]+)")


def log(message: str) -> None:
    sys.stderr.write(message + "\n")
    sys.stderr.flush()


def send_message(payload: dict[str, Any]) -> None:
    sys.stdout.write(json.dumps(payload, ensure_ascii=False) + "\n")
    sys.stdout.flush()


def empty_result() -> dict[str, Any]:
    return {
        "basic": {"name": None, "email": None, "phone": None, "location": None},
        "education": [],
        "experience": [],
        "projects": [],
        "skills": [],
        "raw_text": "",
    }


def normalize_title(line: str) -> str:
    text = line.strip()
    text = re.sub(r"^[\d\.、\s]+", "", text)
    text = re.sub(r"[:：]+$", "", text)
    return text.strip()


def identify_section(line: str) -> str | None:
    normalized = normalize_title(line)
    for section, keywords in SECTION_KEYWORDS.items():
        if normalized in keywords:
            return section
    if normalized in SHORT_SECTION_KEYWORDS:
        return SHORT_SECTION_KEYWORDS[normalized]
    return None


def chars_to_line_text(chars: list[dict[str, Any]]) -> str:
    return "".join(c["text"] for c in sorted(chars, key=lambda item: item["x0"]))


def group_chars_into_lines(chars: list[dict[str, Any]], tolerance: float = LINE_TOP_TOLERANCE) -> list[str]:
    if not chars:
        return []

    ordered = sorted(chars, key=lambda item: item["top"])
    line_groups: list[list[dict[str, Any]]] = []
    current_group: list[dict[str, Any]] = []
    anchor_top: float | None = None

    for char in ordered:
        top = float(char["top"])
        if anchor_top is None or abs(top - anchor_top) <= tolerance:
            current_group.append(char)
            if anchor_top is None:
                anchor_top = top
        else:
            if current_group:
                line_groups.append(current_group)
            current_group = [char]
            anchor_top = top

    if current_group:
        line_groups.append(current_group)

    return [chars_to_line_text(group) for group in line_groups]


def find_column_split(chars: list[dict[str, Any]], page_width: float) -> float | None:
    if not chars or page_width <= 0:
        return None

    xs = sorted({round(float(char["x0"]), 1) for char in chars})
    if len(xs) < 2:
        return None

    mid_start = page_width * 0.35
    mid_end = page_width * 0.65
    best_gap = 0.0
    split_at: float | None = None

    for left, right in zip(xs, xs[1:]):
        gap = right - left
        center = (left + right) / 2
        if mid_start <= center <= mid_end and gap > best_gap:
            best_gap = gap
            split_at = center

    if split_at is not None and best_gap >= 20:
        return split_at
    return None


def extract_pdf_lines(pdf_path: Path) -> tuple[list[str], int, int]:
    all_lines: list[str] = []
    page_count = 0
    column_count = 1

    with pdfplumber.open(pdf_path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            chars = page.chars or []
            if not chars:
                continue

            split = find_column_split(chars, float(page.width))
            if split is None:
                all_lines.extend(group_chars_into_lines(chars))
                continue

            column_count = 2
            left_chars = [char for char in chars if float(char["x0"]) < split]
            right_chars = [char for char in chars if float(char["x0"]) >= split]
            all_lines.extend(group_chars_into_lines(left_chars))
            all_lines.extend(group_chars_into_lines(right_chars))

    return all_lines, page_count, column_count


def paragraph_heading_level(paragraph: Any) -> int | None:
    style_name = paragraph.style.name if paragraph.style else None
    if not style_name:
        return None
    match = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
    if match:
        return int(match.group(1))
    if style_name in ("Title", "标题"):
        return 1
    return None


def extract_docx_lines(docx_path: Path) -> list[str]:
    document = Document(str(docx_path))
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    return lines


def parse_docx_projects(docx_path: Path) -> list[dict[str, Any]]:
    document = Document(str(docx_path))
    in_projects = False
    projects: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        section = identify_section(text)
        if section == "projects":
            in_projects = True
            current = None
            continue
        if in_projects and section is not None:
            break
        if not in_projects:
            continue

        level = paragraph_heading_level(paragraph)
        if level is not None and level >= 2:
            if current is not None:
                projects.append(current)
            current = {"name": text, "role": None, "highlights": []}
            continue

        if text.startswith(("-", "•", "·", "*")):
            highlight = re.sub(r"^[-•·*]\s*", "", text).strip()
            if current is not None and highlight:
                current["highlights"].append(highlight)
            continue

        if current is not None:
            if current.get("role") is None:
                current["role"] = text
            else:
                current["highlights"].append(text)

    if current is not None:
        projects.append(current)

    return projects


def extract_name(full_text: str) -> str | None:
    match = NAME_RE.search(full_text)
    if match:
        name = match.group("name").strip()
        return name or None
    return None


def extract_phone(full_text: str) -> str | None:
    match = PHONE_RE.search(full_text)
    return match.group(0) if match else None


def extract_email(full_text: str, phone: str | None) -> str | None:
    match = EMAIL_RE.search(full_text)
    if not match:
        return None

    email = match.group(0)
    local, _, domain = email.partition("@")
    if phone and local.startswith(phone):
        local = local[len(phone) :]
        if not local:
            return None
        email = f"{local}@{domain}"
    return email


def extract_location(lines: list[str]) -> str | None:
    for line in lines:
        match = re.search(r"(?:地址|所在地|Location)\s*[：:]\s*(?P<location>.+)", line, re.IGNORECASE)
        if match:
            return match.group("location").strip()
    return None


def split_sections(lines: list[str]) -> tuple[dict[str, list[str]], list[str]]:
    sections: dict[str, list[str]] = {
        "education": [],
        "experience": [],
        "projects": [],
        "skills": [],
    }
    preamble: list[str] = []
    current: str | None = None

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        section = identify_section(stripped)
        if section is not None:
            current = section
            continue

        if current is None:
            preamble.append(stripped)
        else:
            sections[current].append(stripped)

    return sections, preamble


def parse_education_lines(lines: list[str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for line in lines:
        date_match = DATE_RANGE_RE.search(line)
        if date_match:
            if current is None:
                current = {"school": None, "degree": None, "major": None, "start": None, "end": None}
            current["start"] = date_match.group("start")
            current["end"] = date_match.group("end")
            before = line[: date_match.start()].strip(" |，,;；")
            if before and current.get("school") is None:
                current["school"] = before
            entries.append(current)
            current = None
            continue

        major_match = MAJOR_INLINE_RE.search(line)
        if major_match:
            if current is None:
                current = {"school": None, "degree": None, "major": None, "start": None, "end": None}
            current["major"] = major_match.group("major")
            continue

        if line in ("本科", "硕士", "博士", "大专", "专科", "高中"):
            if current is None:
                current = {"school": None, "degree": None, "major": None, "start": None, "end": None}
            current["degree"] = line
            continue

        if current is None:
            current = {"school": line, "degree": None, "major": None, "start": None, "end": None}
        elif current.get("school") is None:
            current["school"] = line
        elif current.get("degree") is None and line in ("本科", "硕士", "博士", "大专"):
            current["degree"] = line
        else:
            entries.append(current)
            current = {"school": line, "degree": None, "major": None, "start": None, "end": None}

    if current is not None:
        entries.append(current)

    return [entry for entry in entries if any(entry.values())]


def parse_experience_lines(lines: list[str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for line in lines:
        date_match = DATE_RANGE_RE.search(line)
        if date_match:
            start = date_match.group("start")
            end = date_match.group("end")
            before = line[: date_match.start()].strip(" |，,;；")
            company = None
            title = None

            if before:
                parts = [part.strip() for part in re.split(r"\s*\|\s*", before) if part.strip()]
                if len(parts) >= 2:
                    company, title = parts[0], parts[1]
                elif len(parts) == 1:
                    company = parts[0]

            current = {
                "company": company,
                "title": title,
                "start": start,
                "end": end,
                "highlights": [],
            }
            entries.append(current)
            continue

        if line.startswith(("-", "•", "·", "*")):
            highlight = re.sub(r"^[-•·*]\s*", "", line).strip()
            if highlight and current is not None:
                current["highlights"].append(highlight)

    return entries


def parse_project_lines(lines: list[str]) -> list[dict[str, Any]]:
    projects: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for line in lines:
        if line.startswith(("-", "•", "·", "*")):
            detail = re.sub(r"^[-•·*]\s*", "", line).strip()
            if current is not None and detail:
                current.setdefault("highlights", []).append(detail)
            continue

        if current is not None:
            if current.get("role") is None and not DATE_RANGE_RE.search(line):
                current["role"] = line
                continue
            projects.append(current)

        current = {"name": line, "role": None, "highlights": []}

    if current is not None:
        projects.append(current)

    return projects


def parse_skill_lines(lines: list[str]) -> list[str]:
    skills: list[str] = []
    for line in lines:
        parts = re.split(r"[,，、/|;\s]+", line)
        for part in parts:
            skill = part.strip()
            if skill and skill not in skills:
                skills.append(skill)
    return skills


def parse_inline_education(full_text: str) -> list[dict[str, Any]]:
    match = DEGREE_INLINE_RE.search(full_text)
    if not match:
        return []

    entry: dict[str, Any] = {
        "school": None,
        "degree": match.group("degree"),
        "major": match.group("major"),
        "start": None,
        "end": None,
    }
    if entry["major"] is None:
        major_match = MAJOR_INLINE_RE.search(full_text)
        if major_match:
            entry["major"] = major_match.group("major")
    return [entry]


def parse_resume(file_path: str) -> dict[str, Any]:
    """将简历 PDF/DOCX 解析为固定 schema 的结构化 JSON（无状态，每次独立解析）。"""

    # 把传入路径转为绝对路径，避免相对路径歧义
    path = Path(file_path).resolve()
    # 允许的根目录：仅 data/resumes/，禁止读取其它目录
    resumes_root = RESUMES_ROOT
    if not str(path).startswith(str(resumes_root)):
        raise ValueError("file_path must be under data/resumes")

    # 文件必须真实存在，否则无法打开
    if not path.is_file():
        raise ValueError(f"file not found: {path}")

    # 按扩展名选择解析器；统一转小写便于比较
    suffix = path.suffix.lower()
    column_count = 1  # 默认单栏，PDF 分栏检测后可能变为 2
    if suffix == ".pdf":
        # PDF：按字符坐标收行，支持双栏；返回行列表、页数、栏数
        lines, page_count, column_count = extract_pdf_lines(path)
        log(f"parse_resume open {suffix.lstrip('.')} pages={page_count}")
    elif suffix == ".docx":
        # DOCX：按段落提取纯文本行；项目段另走样式层级解析
        lines = extract_docx_lines(path)
        page_count = 1  # docx 无固定页概念，日志统一记 1
        docx_projects = parse_docx_projects(path)
        log(f"parse_resume open {suffix.lstrip('.')} pages={page_count}")
    else:
        # 非 pdf/docx 直接拒绝，不做猜测
        raise ValueError(f"unsupported file type: {suffix}")

    # 记录分栏结果，便于排查 PDF 双栏排版问题
    log(f"parse_resume columns={column_count}")

    # 按章节标题切分：preamble=标题前的基本信息行，sections=各章节正文行
    sections, preamble = split_sections(lines)
    # 合并全文（不含章节标题行），供姓名/邮箱/电话等全文检索
    all_content_lines = preamble + [line for group in sections.values() for line in group]
    # raw_text 保留原始阅读顺序的全部非空行，供审计与兜底
    raw_text = "\n".join(line for line in lines if line.strip())

    # 先抽手机号，后续邮箱 local-part 去重时需要用到
    full_text = "\n".join(all_content_lines)
    phone = extract_phone(full_text)
    # 初始化固定输出 schema，再逐字段填充
    result = empty_result()
    result["basic"]["name"] = extract_name(full_text)  # 全文搜「姓名：」「Name:」
    result["basic"]["phone"] = phone
    result["basic"]["email"] = extract_email(full_text, phone)  # 去掉邮箱前缀里的手机号
    result["basic"]["location"] = extract_location(preamble + all_content_lines)
    result["raw_text"] = raw_text

    # 教育：优先解析「教育背景」章节；无章节时回退到行内「学历：本科」等
    education = parse_education_lines(sections["education"])
    if not education:
        education = parse_inline_education(full_text)
    result["education"] = education
    # 工作经历：只取「工作经历」标题与下一标题之间的行，一条任职一行
    result["experience"] = parse_experience_lines(sections["experience"])
    # 项目：docx 用段落 Heading 层级；pdf 或 docx 解析失败时回退行级规则
    if suffix == ".docx":
        result["projects"] = docx_projects or parse_project_lines(sections["projects"])
    else:
        result["projects"] = parse_project_lines(sections["projects"])
    # 技能：只取「技能/技术专长」等标题下内容，不混入工作经历
    result["skills"] = parse_skill_lines(sections["skills"])

    # 打阶段日志：只记计数，不输出正文或完整 JSON
    log(
        "parse_resume sections "
        f"education={len(result['education'])} "
        f"experience={len(result['experience'])} "
        f"projects={len(result['projects'])} "
        f"skills={len(result['skills'])}"
    )
    return result


def handle_initialize(request_id: Any, _params: dict[str, Any]) -> None:
    send_message(
        {
            "jsonrpc": "2.0",
            "id": request_id,
            "result": {
                "protocolVersion": "2025-03-26",
                "capabilities": {"tools": {}},
                "serverInfo": {"name": "resume-parser", "version": "0.1.0"},
            },
        }
    )


def handle_tools_list(request_id: Any) -> None:
    send_message(
        {
            "jsonrpc": "2.0",
            "id": request_id,
            "result": {
                "tools": [
                    {
                        "name": "parse_resume",
                        "description": "Parse a resume PDF or DOCX file into structured fields.",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "file_path": {
                                    "type": "string",
                                    "description": "Absolute path to a resume file under data/resumes",
                                }
                            },
                            "required": ["file_path"],
                        },
                    }
                ]
            },
        }
    )


def handle_tools_call(request_id: Any, params: dict[str, Any]) -> None:
    tool_name = params.get("name", "")
    log(f"tools/call {tool_name} begin")
    started = time.monotonic()

    try:
        if tool_name != "parse_resume":
            raise ValueError(f"unknown tool: {tool_name}")

        arguments = params.get("arguments") or {}
        file_path = arguments.get("file_path")
        if not file_path:
            raise ValueError("file_path is required")

        parsed = parse_resume(file_path)
        elapsed_ms = int((time.monotonic() - started) * 1000)
        log(f"tools/call {tool_name} end {elapsed_ms}ms")
        send_message(
            {
                "jsonrpc": "2.0",
                "id": request_id,
                "result": {
                    "content": [
                        {
                            "type": "text",
                            "text": json.dumps(parsed, ensure_ascii=False),
                        }
                    ],
                    "isError": False,
                },
            }
        )
    except Exception as exc:
        elapsed_ms = int((time.monotonic() - started) * 1000)
        log(f"tools/call {tool_name} end {elapsed_ms}ms error={exc}")
        send_message(
            {
                "jsonrpc": "2.0",
                "id": request_id,
                "error": {
                    "code": -32000,
                    "message": str(exc),
                },
            }
        )


def handle_request(message: dict[str, Any]) -> None:
    method = message.get("method")
    request_id = message.get("id")
    params = message.get("params") or {}

    if method == "initialize":
        handle_initialize(request_id, params)
        return

    if method == "notifications/initialized":
        return

    if method == "tools/list":
        handle_tools_list(request_id)
        return

    if method == "tools/call":
        handle_tools_call(request_id, params)
        return

    if method is not None:
        send_message(
            {
                "jsonrpc": "2.0",
                "id": request_id,
                "error": {
                    "code": -32601,
                    "message": f"Method not found: {method}",
                },
            }
        )


def main() -> None:
    log("resume-parser ready")
    for raw_line in sys.stdin:
        line = raw_line.strip()
        if not line:
            continue
        try:
            message = json.loads(line)
        except json.JSONDecodeError:
            continue
        if not isinstance(message, dict):
            continue
        handle_request(message)


if __name__ == "__main__":
    main()
